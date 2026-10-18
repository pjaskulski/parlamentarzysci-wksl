"""Eksport kart do DOCX w pamięci; dane pochodzą z tych samych kontekstów co PDF."""
from io import BytesIO
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def document(title, landscape=False):
    doc=Document(); section=doc.sections[0]
    section.page_width=Mm(297 if landscape else 210)
    section.page_height=Mm(210 if landscape else 297)
    section.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.top_margin=section.bottom_margin=Mm(20)
    section.left_margin=section.right_margin=Mm(20)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.name='DejaVu Sans';style.font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(23)
    for border in list(doc.styles['Title'].element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.styles['Heading 1'].font.size=Pt(15)
    doc.core_properties.title=title
    doc.core_properties.author='Pracownia Historii Cyfrowej IHPAN'
    footer=section.footer.paragraphs[0]
    footer.add_run('Pracownia Historii Cyfrowej IHPAN, 2026 · ').font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    return doc


def table(doc, headers, rows, widths):
    if doc.paragraphs:doc.paragraphs[-1].paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    available=doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
    for col,width in zip(t.columns,widths):col.width=int(available*width)
    for cell,width in zip(t.rows[0].cells,widths):cell.width=int(available*width)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:b.set(qn('w:'+key),value)
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for cell,text in zip(t.rows[0].cells,headers):
        cell.text=str(text)
        for run in cell.paragraphs[0].runs:run.bold=True
        shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'EFEDE7');cell._tc.get_or_add_tcPr().append(shading)
    for values in rows:
        cells=t.add_row().cells
        for cell,value,width in zip(cells,values,widths):
            cell.width=int(available*width);cell.text=str(value or '-')
    for row in t.rows:
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
                for run in p.runs:run.font.size=Pt(9)
    if not any(headers):
        t._tbl.remove(t.rows[0]._tr)
    else:
        for cell in t.rows[0].cells:
            for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def finish(doc):
    stream=BytesIO();doc.save(stream);stream.seek(0);return stream


def build_person_docx(context, *, app_title, format_source, certainty):
    p=context['person'];record=context['record'];doc=document(p['name'])
    doc.add_paragraph(app_title);doc.add_paragraph(p['name'],'Title')
    dates=[label+': '+p[key] for key,label in [('birth_date_raw','Data urodzenia'),('baptism_date_raw','Data chrztu'),('death_date_raw','Data śmierci')] if p.get(key)]
    doc.add_paragraph(' · '.join(dates) or 'Daty życia niepodane w źródle')
    names=[]
    for key,label in [('given_names','Imiona'),('surname','Nazwisko'),('patronymic','Otczestwo'),('surname_variants','Warianty nazwiska'),('given_name_variants','Warianty imienia'),('patronymic_variants','Warianty otczestwa'),('other_name_variants','Pozostałe człony nazwy')]:
        value=p.get(key)
        if value:names.append((label,'; '.join(value) if isinstance(value,list) else value))
    if names:doc.add_heading('Imiona / Nazwiska / Warianty',1);table(doc,['',''],names,[.3,.7])
    doc.add_heading('Tekst biogramu',1);doc.add_paragraph(p['biogram'])
    if context['offices']:
        doc.add_heading('Urzędy',1);rows=[]
        for o in context['offices']:
            name=o['name']
            if o['office_type']:name+='\nRodzaj: '+o['office_type']
            if o['certainty']=='odrzucone_lub_prostowane' and not o['name'].lower().startswith('niedoszły '):name+='\n'+certainty.get(o['certainty'],o['certainty'] or '')
            rows.append((name,o['dating_raw'] or 'Niepodane'))
        table(doc,['Urząd','Datowanie'],rows,[.73,.27])
    if context['titles']:
        doc.add_heading('Inne stopnie, godności i określenia rodzinne',1)
        for t in context['titles']:
            label=doc.add_paragraph()
            label.paragraph_format.space_after=Pt(2)
            run=label.add_run('Określenie rodzinne' if t['kind']=='family_designation' else 'Stopień lub godność')
            run.font.size=Pt(8)
            run.font.color.rgb=RGBColor.from_string('6C6A63')
            doc.add_paragraph(t['description_raw'])
    if context['roles']:
        doc.add_heading('Funkcje parlamentarne',1)
        doc.add_paragraph('/I, /II, /III oznacza kolejność w roku, /k — konwokację. Znak ? wskazuje niepewność.')
        rows=[]
        for r in context['roles']:
            event=r['event_code_raw']+'\nWydarzenie nieustalone'
            if r['event_id']:
                event=r['event_name']+' ('+r['event_code_raw']+')\n'+(r['event_place'] or 'Miejsce niepodane')+'\n'+r['event_date_raw']
                if r['occurrence_status'] in ('not_held','cancelled'):event+='\nObrady nie odbyły się'
            role=r['role'].capitalize()+(' — mandat nieuznany' if r['attendance']=='mandat_nieuznany' else '')
            rows.append((event,role,'; '.join(r.get('representations',[])) or '-'))
        table(doc,['Sejm lub zjazd','Rola','Posłował z'],rows,[.52,.14,.34])
    def items(values,bibliography=False):
        for item in values:
            text=item.get('zapis',item.get('cel','')) if isinstance(item,dict) else item
            doc.add_paragraph(format_source(text) if bibliography else str(text))
            if isinstance(item,dict):
                note=item.get('uwaga');level=item.get('pewność')
                if note=='Fragment listy kariery; możliwa elipsa urzędu lub kontynuacja datowania.':note=None
                label=certainty.get(level,level or '') if level!='stwierdzone_w_źródle' else ''
                if level=='niepewne' and str(text).lower().startswith('prawdopodobnie '):label=''
                notes=[label,note]
                if any(notes):doc.add_paragraph(' · '.join(str(x) for x in notes if x))
    for key,values in context['sections']:
        if values:doc.add_heading(key.replace('_',' ').capitalize(),1);items(values,key=='źródła')
    return finish(doc)


def build_event_docx(context, *, app_title, role_names, translate=lambda text:text):
    e=context['event'];doc=document(translate(e['name']),True)
    doc.add_paragraph(app_title)
    if e['codes']:doc.add_paragraph(' · '.join(e['codes']))
    doc.add_paragraph(translate(e['name']),'Title')
    doc.add_paragraph(' · '.join(filter(None,[e.get('place'),e.get('date_raw')])))
    doc.add_paragraph(translate('Osoby:')+' '+str(context['people_count']))
    if e['occurrence_status'] in ('not_held','cancelled'):doc.add_paragraph(translate('Obrady nie odbyły się.'))
    if len(e['sessions'])>1:
        doc.add_heading(translate('Części obrad'),1)
        for s in e['sessions']:
            text=' · '.join(filter(None,[s.get('place') or e.get('place'),s.get('date_raw')]))
            if s.get('continuation'):text+=' · '+translate('Kontynuacja po limicie')
            doc.add_paragraph(text)
    footnotes=[]
    def references(notes):
        refs=[]
        for n in notes:
            key=(n['label'],n['explanation'])
            if key not in footnotes:footnotes.append(key)
            num=footnotes.index(key)+1
            if num not in refs:refs.append(num)
        return ' '+''.join(f'[{n}]' for n in refs) if refs else ''
    if not context['groups']:doc.add_paragraph(translate('Brak powiązanych osób.'))
    for role,people in context['groups'].items():
        doc.add_heading(translate(role_names.get(role,role.capitalize()))+' ('+str(len(people))+')',1)
        headers=['Osoba','Lata życia']+(['Posłował z'] if role=='poseł' else [])+['Urzędy'];rows=[]
        for p in people:
            life=p['life_years']
            if translate('Lata życia')!='Lata życia':life=life.replace('ur.','b.').replace('zm.','d.')
            rows.append([p['name']+references(p['participation_notes']),life]+(['; '.join(p['representations']) or '-'] if role=='poseł' else [])+[', '.join(p['offices']) or '-'])
        table(doc,[translate(h) for h in headers],rows,[.27,.13,.17,.43] if role=='poseł' else [.29,.14,.57])
    if footnotes:
        doc.add_heading(translate('Zastrzeżenia dotyczące udziału'),1)
        for i,(label,explanation) in enumerate(footnotes,1):doc.add_paragraph(f'[{i}] '+translate(label)+': '+translate(explanation))
    return finish(doc)
