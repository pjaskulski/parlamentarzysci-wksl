"""Eksport wydarzenia i grup uczestników do XLSX, wyłącznie w pamięci."""
from io import BytesIO
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def append(sheet,values):
    sheet.append(values)
    # Dane historyczne są tekstem, także gdy zaczynają się znakiem „=”.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value,str):cell.data_type='s'

def style(sheet,widths):
    sheet.sheet_view.showGridLines=False
    sheet.freeze_panes='A2'
    for col,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(col)].width=width
    for row in sheet:
        lines=1
        for cell in row:
            cell.font=Font(name='Calibri',size=11,color='292723')
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if row[0].row==1:
                cell.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor='62513B')
            elif row[0].row%2==0:cell.fill=PatternFill('solid',fgColor='F5F3EE')
            width=widths[cell.column-1]
            lines=max(lines,sum(max(1,math.ceil(len(part)/max(1,width-3))) for part in str(cell.value or '').split('\n')))
        sheet.row_dimensions[row[0].row].height=min(409,lines*16+10)
    sheet.print_title_rows='1:1'
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='landscape'
    sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.oddFooter.center.text='Pracownia Historii Cyfrowej IHPAN, 2026'
    sheet.oddFooter.right.text='&P / &N'


def build_event_xlsx(context, *, app_title, role_names, translate=lambda text:text):
    workbook=Workbook()
    workbook.properties.creator='Pracownia Historii Cyfrowej IHPAN'
    event=context['event']
    workbook.properties.title=translate(event['name'])

    overview=workbook.active;overview.title=translate('Wydarzenie')
    append(overview,[translate('Informacja'),translate('Treść')])
    for label,value in [('Baza',app_title),('Sejm lub zjazd',translate(event['name'])),
                        ('Kod',' · '.join(event['codes']) or '-'),('Miejsce',event.get('place') or '-'),
                        ('Datowanie',event.get('date_raw') or '-'),('Osoby',context['people_count'])]:
        append(overview,[translate(label),value])
    if event['occurrence_status'] in ('not_held','cancelled'):
        append(overview,[translate('Uwagi'),translate('Obrady nie odbyły się.')])
    if len(event['sessions'])>1:
        for session in event['sessions']:
            text=' · '.join(filter(None,[session.get('place') or event.get('place'),session.get('date_raw')]))
            if session.get('continuation'):text+=' · '+translate('Kontynuacja po limicie')
            append(overview,[translate('Części obrad'),text])
    if not context['groups']:append(overview,[translate('Uwagi'),translate('Brak powiązanych osób.')])
    style(overview,[26,100])
    for role,people in context['groups'].items():
        name=re.sub(r'[\\/*?:\[\]]',' ',translate(role_names.get(role,role.capitalize())))[:31] or 'Osoby'
        sheet=workbook.create_sheet(name)
        headers=['Osoba','Lata życia']+(['Posłował z'] if role=='poseł' else [])+['Urzędy','Uwagi']
        append(sheet,[translate(h) for h in headers])
        for person in people:
            life=person['life_years']
            if translate('Lata życia')!='Lata życia':life=life.replace('ur.','b.').replace('zm.','d.')
            notes=[]
            for note in person['participation_notes']:
                text=translate(note['explanation'])
                if text not in notes:notes.append(text)
            values=[person['name'],life]+(['; '.join(person['representations']) or '-'] if role=='poseł' else [])
            values+=[', '.join(person['offices']) or '-','\n'.join(notes) or '-']
            append(sheet,values)
        style(sheet,[38,22,28,65,65] if role=='poseł' else [38,22,65,65])
    stream=BytesIO();workbook.save(stream);stream.seek(0);return stream


def build_people_xlsx(people, *, translate=lambda text:text):
    """Cztery podstawowe kolumny; jedna osoba w wierszu."""
    workbook=Workbook();sheet=workbook.active
    sheet.title=translate('Wyniki wyszukiwania')
    workbook.properties.creator='Pracownia Historii Cyfrowej IHPAN'
    workbook.properties.title=translate('Wyniki wyszukiwania')
    append(sheet,[translate(label) for label in ['Osoba','Lata życia','Urzędy','Sejmy i zjazdy']])
    for person in people:
        life=person['life_years']
        if translate('Lata życia')!='Lata życia':life=life.replace('ur.','b.').replace('zm.','d.')
        append(sheet,[person['name'],life,', '.join(person['offices']) or '-',', '.join(person['event_codes']) or '-'])
    style(sheet,[42,24,75,75])
    stream=BytesIO();workbook.save(stream);stream.seek(0);return stream


def build_events_xlsx(events, event_codes, *, translate=lambda text:text):
    """Kolumny i kolejność zgodne z listą wydarzeń w katalogu."""
    workbook=Workbook();sheet=workbook.active
    sheet.title=translate('Sejmy i zjazdy')
    workbook.properties.creator='Pracownia Historii Cyfrowej IHPAN'
    workbook.properties.title=translate('Sejmy i zjazdy')
    append(sheet,[translate(label) for label in ['Kod','Sejm lub zjazd','Miejsce','Data','Osoby']])
    for event in events:
        append(sheet,[' · '.join(event_codes.get(event['event_id'],[])) or '—',
                      translate(event['name']),event['place'] or translate('Miejsce niepodane w nagłówku'),
                      event['date_raw'],event['people_count']])
    style(sheet,[26,62,28,68,12])
    stream=BytesIO();workbook.save(stream);stream.seek(0);return stream
